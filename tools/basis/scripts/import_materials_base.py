"""Импорт производственной базы материалов БАЗИС (xlsx) → materials/baza_materiala.json.

База — реальный справочник производства (≈5000 позиций): листовой материал,
кромка, погонаж, крепёж, фурнитура. Нормализуем в машинный JSON со стабильными
id, чтобы система выбирала РЕАЛЬНЫЕ материалы/фурнитуру (артикул + имя БАЗИС),
а не выдумывала их по тексту ТЗ (AKD-11).

Запуск:
  python scripts/import_materials_base.py [path/to/База материала.xlsx]
По умолчанию берёт materials/source/baza_materiala.xlsx.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SRC = ROOT / "materials" / "source" / "baza_materiala.xlsx"
OUT = ROOT / "materials" / "baza_materiala.json"

# Русский заголовок → чистый ключ
COLS = {
    "Артикул материала": "article",
    "Наименование материала": "name",
    "Номер группы": "group_id",
    "Наименование группы": "group",
    "Единица измерения": "unit",
    "Стоимость": "cost",
    "Коэффициент": "coef",
    "Длина": "length",
    "Ширина": "width",
    "Толщина": "thickness",
    "Обозначение": "designation",     # кандидат точного имени/кода БАЗИС
    "Свес": "overhang",
    "Цвет (целое число)": "color_code",
    "Текстура": "texture",
    "Класс": "cls",
    "Тип материала": "material_type",
    "Идентификатор для синхронизации": "sync_id",  # id для синхронизации с БАЗИС
    "Масса": "mass",
    "Комментарий": "comment",
}
NUMERIC = {"cost", "coef", "length", "width", "thickness", "overhang", "mass"}

# Верхняя категория по числовому префиксу группы ("01 Листовой материал/...")
CATEGORY = {
    "01": "Листовой материал",
    "02": "Кромочные материалы",
    "03": "Погонные материалы",
    "04": "Крепёж",
    "05": "Фурнитура",
    "06": "Прочее",
    "07": "Прочее",
    "08": "Виртуальные материалы",
}


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip().replace(",", ".")
    try:
        f = float(s)
        return int(f) if f == int(f) else f
    except ValueError:
        return str(v).strip()


def _category(group: str) -> str:
    g = (group or "").strip()
    head = g.split(" ", 1)[0].split("/", 1)[0]
    if head in CATEGORY:
        return CATEGORY[head]
    return g.split("/", 1)[0].strip() or "Прочее"


def convert(src: Path) -> dict:
    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    keys = [COLS.get(h, None) for h in header]

    items: list[dict] = []
    cat_counts: dict[str, int] = {}
    group_counts: dict[str, int] = {}
    for i, row in enumerate(rows):
        rec: dict = {}
        for k, val in zip(keys, row):
            if k is None or val in (None, ""):
                continue
            rec[k] = _num(val) if k in NUMERIC else (str(val).strip() if isinstance(val, str) else val)
        if not rec.get("name") and not rec.get("article"):
            continue
        rec["category"] = _category(rec.get("group", ""))
        # стабильный id: sync_id → article → порядковый
        sid = rec.get("sync_id") or rec.get("article")
        rec["id"] = str(sid) if sid not in (None, "") else f"row{i}"
        items.append(rec)
        cat_counts[rec["category"]] = cat_counts.get(rec["category"], 0) + 1
        if rec.get("group"):
            group_counts[rec["group"]] = group_counts.get(rec["group"], 0) + 1

    return {
        "schemaVersion": "material-base-v1",
        "source": "База материала.xlsx — производственный справочник (≈5000 позиций)",
        "note": "Нормализованная база материалов/фурнитуры. id = sync_id (для синхронизации с БАЗИС) или артикул. "
                "designation/name — кандидаты точного имени в БАЗИС. Используется materials.py для подбора реальных позиций.",
        "count": len(items),
        "categories": dict(sorted(cat_counts.items(), key=lambda kv: -kv[1])),
        "groups": dict(sorted(group_counts.items(), key=lambda kv: -kv[1])),
        "items": items,
    }


def main() -> int:
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SRC
    if not src.exists():
        print(f"Не найден исходник: {src}", file=sys.stderr)
        return 2
    data = convert(src)
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"Готово: {OUT}  ({data['count']} позиций, {len(data['categories'])} категорий, {len(data['groups'])} групп)")
    for c, n in data["categories"].items():
        print(f"  {n:>5}  {c}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
